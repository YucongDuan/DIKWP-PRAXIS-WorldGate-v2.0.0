from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
pages=json.loads((ROOT/'docs/report_pages.json').read_text())
NAVY='15374A';TEAL='126B73';GRAY='56636C';LIGHT='EEF4F5'
def font(run,size=None,bold=None,color=None,mono=False):
    run.font.name='DejaVu Sans Mono' if mono else 'DejaVu Sans'
    pr=run._element.get_or_add_rPr();rf=pr.rFonts
    if rf is None:rf=OxmlElement('w:rFonts');pr.append(rf)
    rf.set(qn('w:ascii'),'DejaVu Sans Mono' if mono else 'DejaVu Sans');rf.set(qn('w:hAnsi'),'DejaVu Sans Mono' if mono else 'DejaVu Sans');rf.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    if size:run.font.size=Pt(size)
    if bold is not None:run.font.bold=bold
    if color:run.font.color.rgb=RGBColor.from_string(color)
def inline(p,text,size=10.4,color=None):
    for i,s in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)',text)):
        bold=s.startswith('**');mono=s.startswith('`')
        if bold:s=s[2:-2]
        elif mono:s=s[1:-1]
        r=p.add_run(s);font(r,size,bold,color,mono)
    return p
def shade(elem,fill):
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);elem.append(sh)
for lang in ('cn','en'):
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929);sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7);sec.header_distance=Inches(.25);sec.footer_distance=Inches(.28)
    normal=d.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10.4);normal.paragraph_format.line_spacing=1.17;normal.paragraph_format.space_after=Pt(6)
    for st,size in [('Title',25),('Subtitle',13),('Heading 1',18),('Heading 2',12),('Heading 3',11)]:
        s=d.styles[st];s.font.name='DejaVu Sans';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(NAVY if st in ('Title','Heading 1') else TEAL);s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(7)
    h=sec.header.paragraphs[0];inline(h,'DIKWP / PRAXIS     WORLDGATE 2.0.0',8.3,GRAY)
    pp=h._p.get_or_add_pPr();bd=OxmlElement('w:pBdr');bt=OxmlElement('w:bottom');bt.set(qn('w:val'),'single');bt.set(qn('w:sz'),'6');bt.set(qn('w:color'),TEAL);bd.append(bt);pp.append(bd)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    inline(f,('本地工程参考 · ' if lang=='cn' else 'Local engineering reference · '),8.2,GRAY)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    for pi,page in enumerate(pages):
        if pi:d.add_page_break()
        lines=page[lang].splitlines();i=0
        while i<len(lines):
            line=lines[i]
            if not line.strip():i+=1;continue
            if line.startswith('```'):
                code=[];i+=1
                while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
                p=d.add_paragraph();p.paragraph_format.left_indent=Inches(.10);p.paragraph_format.right_indent=Inches(.06);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.05
                shade(p._p.get_or_add_pPr(),LIGHT)
                r=p.add_run('\n'.join(code));font(r,8.1,False,NAVY,True)
                p.paragraph_format.keep_together=True;i+=1;continue
            if line.startswith('|'):
                rows=[]
                while i<len(lines) and lines[i].startswith('|'):
                    ss=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                    if not all(re.fullmatch(r':?-+:?',x) for x in ss):rows.append(ss)
                    i+=1
                t=d.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
                total=6.84;n=len(rows[0]);widths=([2.15,4.69] if n==2 else ([1.28,2.5,3.06] if n==3 else [total/n]*n))
                for col,w in zip(t.columns,widths):col.width=Inches(w)
                for ri,row in enumerate(rows):
                    cells=t.add_row().cells
                    for j,(c,txt) in enumerate(zip(cells,row)):
                        c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        tcPr=c._tc.get_or_add_tcPr();shade(tcPr,NAVY if ri==0 else ('F1F5F6' if ri%2 else 'FFFFFF'))
                        mar=OxmlElement('w:tcMar')
                        for side in ('top','left','bottom','right'):
                            e=OxmlElement('w:'+side);e.set(qn('w:w'),'70' if side in ('top','bottom') else '85');e.set(qn('w:type'),'dxa');mar.append(e)
                        tcPr.append(mar)
                        p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.07
                        inline(p,txt,9.0,'FFFFFF' if ri==0 else '1B2830')
                        for r in p.runs:
                            if ri==0:r.bold=True
                    tr=t.rows[-1]._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');tr.append(ns)
                    if ri==0:rh=OxmlElement('w:tblHeader');tr.append(rh)
                d.add_paragraph().paragraph_format.space_after=Pt(2);continue
            if line.startswith('# '):
                p=d.add_paragraph(style='Title' if pi==0 else 'Heading 1');inline(p,line[2:],25 if pi==0 else 18,NAVY)
            elif line.startswith('## '):
                p=d.add_paragraph(style='Subtitle' if pi==0 else 'Heading 2');inline(p,line[3:],14 if pi==0 else 12,TEAL)
            elif line.startswith('### '):
                p=d.add_paragraph();inline(p,line[4:],12,TEAL)
            else:
                p=d.add_paragraph()
                if pi==16 and lang=='en':
                    p.paragraph_format.line_spacing=1.07;p.paragraph_format.space_after=Pt(4);inline(p,line,9.5)
                else:inline(p,line)
            i+=1
    d.core_properties.title='DIKWP-PRAXIS-OS 2.0.0 WORLDGATE — '+('对比升维与工程报告' if lang=='cn' else 'Comparison and Engineering Report');d.core_properties.author='Yucong Duan';d.core_properties.subject='AI-assisted engineering release, bounded local reference implementation';d.core_properties.comments='No external certification or endorsement is asserted.';d.core_properties.language='zh-CN' if lang=='cn' else 'en-US'
    out=ROOT/f'docs/report_{lang}.docx';d.save(out);print(out)
